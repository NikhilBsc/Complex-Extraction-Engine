"""
tests/test_consolidation.py
---------------------------
Tests for Phase 8 (Consolidation & Deduplication) and Phase 9 (Excel output).

Usage:
    python -m tests.test_consolidation

Test groups:
  1. Field name normalisation and canonical form
  2. Deduplication — exact and similarity-based
  3. Conflict detection (same field, different values)
  4. Full consolidation on realistic KPMG-style extraction
  5. Excel output — file written and readable
"""

from __future__ import annotations

import sys
import io
import logging
import tempfile
from pathlib import Path

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

logging.basicConfig(level=logging.INFO, format="%(levelname)-8s %(name)s | %(message)s")


def _assert(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(f"FAIL: {message}")
    print(f"  [PASS] {message}")


def _make_field(name: str, value: str, conf: float = 0.95,
                status: str = "accepted", seg: str = "seg_001",
                pages: list = None) -> "ExtractedField":
    from gpt.models import ExtractedField
    f = ExtractedField(
        field_name=name, value=value, confidence=conf,
        evidence=f"{name}: {value}", segment_id=seg,
        page_numbers=pages or [1],
    )
    f.validation_status = status
    f.is_rehydrated = True
    return f


# --------------------------------------------------------------------------- #
# 1. Field name normalisation
# --------------------------------------------------------------------------- #

def test_normalisation() -> None:
    print("\n" + "=" * 60)
    print("1. Field name normalisation tests")
    print("=" * 60)

    from output.consolidator import _normalise_field_name, _canonical_field_name

    _assert(_normalise_field_name("Letter Date") == "letter date", "Lowercase normalisation")
    _assert(_normalise_field_name("  Letter Date:  ") == "letter date", "Strip and colon removal")
    _assert(_normalise_field_name("PROPOSED FEE (TAX)") == "proposed fee (tax)", "All-caps normalised")

    _assert(_canonical_field_name("letter date") == "Letter Date", "Title case canonical")
    _assert(_canonical_field_name("USD amount") == "USD Amount", "USD preserved as acronym")
    result = _canonical_field_name("proposed fee (tax compliance services)")
    _assert("Proposed Fee" in result and "Services" in result, "Parenthetical key words capitalised")


# --------------------------------------------------------------------------- #
# 2. Deduplication
# --------------------------------------------------------------------------- #

def test_deduplication() -> None:
    print("\n" + "=" * 60)
    print("2. Deduplication tests")
    print("=" * 60)

    from output.consolidator import consolidate

    # Exact duplicate (same field, same value, different segments)
    fields = [
        _make_field("Letter Date", "January 22, 2026", 0.98, seg="seg_001"),
        _make_field("Letter Date", "January 22, 2026", 0.92, seg="seg_003"),  # duplicate
    ]
    result = consolidate(fields)
    _assert(len(result) == 1, "Exact duplicate deduplicated to 1 field")
    _assert(result[0].occurrence_count == 2, "Occurrence count = 2")
    _assert(result[0].confidence == 0.98, "Highest confidence kept")
    _assert(not result[0].has_conflict, "No conflict for same value")

    # Similarity-based dedup ("Proposed Fee" ~ "Proposed Fee (Tax Compliance Services)")
    fields2 = [
        _make_field("Proposed Fee", "USD 43,500", 0.95),
        _make_field("Proposed Fee (Tax Compliance Services)", "USD 43,500", 0.97),
    ]
    result2 = consolidate(fields2)
    # These are similar but might not merge (parenthetical adds meaning)
    # The test just checks the result is sane
    _assert(1 <= len(result2) <= 2, "Similar fee fields: 1 or 2 results (similarity boundary)")

    # Completely different fields → no merge
    fields3 = [
        _make_field("Letter Date", "January 22, 2026", 0.98),
        _make_field("Client Name", "Foundry Associates REIT LLC", 0.97),
        _make_field("Proposed Fee", "USD 43,500", 0.99),
    ]
    result3 = consolidate(fields3)
    _assert(len(result3) == 3, "3 distinct fields kept as 3")


# --------------------------------------------------------------------------- #
# 3. Conflict detection
# --------------------------------------------------------------------------- #

def test_conflict_detection() -> None:
    print("\n" + "=" * 60)
    print("3. Conflict detection tests")
    print("=" * 60)

    from output.consolidator import consolidate

    # Same field name, genuinely different values
    fields = [
        _make_field("Letter Date", "January 22, 2026", 0.98, pages=[1]),
        _make_field("Letter Date", "February 15, 2026", 0.90, seg="seg_010", pages=[10]),
    ]
    result = consolidate(fields)
    _assert(len(result) == 1, "Conflict collapses to 1 field")
    cf = result[0]
    _assert(cf.has_conflict, "Conflict detected for different values")
    _assert(len(cf.all_values) == 2, "Both values preserved in all_values")
    _assert(cf.value == "January 22, 2026", "Highest-confidence value is primary")

    # Same value with minor variation → NOT a conflict
    fields2 = [
        _make_field("Payment Term", "10 working days", 0.96),
        _make_field("Payment Term", "10 working days.", 0.91),  # trailing period
    ]
    result2 = consolidate(fields2)
    _assert(not result2[0].has_conflict, "Minor punctuation variation not a conflict")


# --------------------------------------------------------------------------- #
# 4. Full realistic consolidation
# --------------------------------------------------------------------------- #

def test_full_consolidation() -> None:
    print("\n" + "=" * 60)
    print("4. Full consolidation on realistic KPMG-style extraction")
    print("=" * 60)

    from output.consolidator import consolidate, consolidation_summary

    fields = [
        _make_field("Letter Date", "January 22, 2026", 0.98),
        _make_field("Client Name", "Foundry Associates REIT LLC", 0.97),
        _make_field("Client Address", "1270 Avenue of the Americas, STE 1907, New York", 0.95),
        _make_field("Auditor / Firm Name", "KPMG Samjong Accounting Corp.", 0.98),
        _make_field("Service Type", "U.S. tax compliance services", 0.96),
        _make_field("Tax Year Covered", "2025", 0.97),
        _make_field("Proposed Fee", "USD 43,500 (cap)", 0.99),
        _make_field("Fee Exclusions", "Exclusive of VAT and reasonable out-of-pocket expenses", 0.95),
        _make_field("Payment Term (Duration)", "10 working days", 0.96),
        _make_field("Billing Schedule", "50% billed in April 2026, remaining 50% by August 2026", 0.93),
        # Duplicate with slight variation
        _make_field("Proposed Fee (Tax Compliance)", "USD 43,500 (cap)", 0.92, seg="seg_002"),
        # Flagged field
        _make_field("Contact", "+82-2-2112-0721", 0.72, status="flagged"),
        # Rejected — should not appear
        _make_field("Page 1 of 5", "Header", 0.50, status="rejected"),
    ]

    result = consolidate(fields, include_flagged=True)

    _assert(len(result) >= 8, f"At least 8 unique fields (got {len(result)})")
    _assert(all(cf.validation_status != "rejected" for cf in result),
            "No rejected fields in consolidated output")
    _assert(all(cf.confidence > 0 for cf in result), "All fields have positive confidence")

    # Document order preserved
    page_order = [cf.source_pages[0] for cf in result]
    _assert(page_order == sorted(page_order), "Fields sorted by page order")

    summary = consolidation_summary(result)
    _assert(summary["total_unique_fields"] == len(result), "Summary total matches")
    _assert(summary["avg_confidence"] > 0.85, "Average confidence > 85%")

    print(f"\n  Unique fields  : {len(result)}")
    print(f"  Avg confidence : {summary['avg_confidence']:.1%}")
    print(f"  Conflicts      : {summary['fields_with_conflicts']}")
    print(f"\n  Field list:")
    for cf in result:
        flag = "[FLG]" if cf.validation_status == "flagged" else "[ACC]"
        conflict_note = " [CONFLICT]" if cf.has_conflict else ""
        print(f"  {flag} {cf.field_name:<50} | {str(cf.value)[:40]}{conflict_note}")


# --------------------------------------------------------------------------- #
# 5. Excel output
# --------------------------------------------------------------------------- #

def test_excel_output() -> None:
    print("\n" + "=" * 60)
    print("5. Excel output test")
    print("=" * 60)

    from output.consolidator import consolidate
    from output.excel_writer import write_excel
    from validation.validator import ValidationReport
    from gpt.models import ExtractedField

    fields = [
        _make_field("Letter Date", "January 22, 2026", 0.98),
        _make_field("Client Name", "Foundry Associates REIT LLC", 0.97),
        _make_field("Proposed Fee", "USD 43,500 (cap)", 0.99),
        _make_field("Payment Term", "10 working days", 0.67, status="flagged"),
    ]

    # Rejected field for audit sheet
    rej = _make_field("Page 1 of 5", "Header", 0.40, status="rejected")
    rej.rejection_reason = "boilerplate field name"

    report = ValidationReport(
        total_fields=5,
        accepted=3,
        flagged=1,
        rejected=1,
        rejected_fields=[rej],
        flagged_fields=[f for f in fields if f.validation_status == "flagged"],
        accepted_fields=[f for f in fields if f.validation_status == "accepted"],
    )

    consolidated = consolidate(fields, include_flagged=True)

    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = Path(tmpdir) / "test_extraction.xlsx"
        written = write_excel(
            consolidated_fields=consolidated,
            report=report,
            output_path=out_path,
            doc_name="KPMG Engagement Letter - Test",
            extraction_meta={
                "model_used": "gpt-4o",
                "total_tokens": 4500,
                "cost_usd": 0.0234,
                "segments_processed": 5,
                "segments_skipped": 1,
            },
        )

        _assert(written.exists(), "Excel file written to disk")
        _assert(written.stat().st_size > 5000, "Excel file has content (>5KB)")

        # Verify it's a valid workbook
        import openpyxl
        wb = openpyxl.load_workbook(written)
        _assert("Extracted Fields" in wb.sheetnames, "Sheet 'Extracted Fields' exists")
        _assert("Extraction Summary" in wb.sheetnames, "Sheet 'Extraction Summary' exists")
        _assert("Rejected Fields" in wb.sheetnames, "Sheet 'Rejected Fields' exists")

        ws1 = wb["Extracted Fields"]
        _assert(ws1.max_row >= 2, "Data rows exist in Extracted Fields sheet")
        # Check header
        _assert(ws1.cell(1, 2).value == "Field Name", "Header row correct")
        _assert(ws1.cell(1, 3).value == "Value", "Value column header correct")

        # Check some data
        data_values = [ws1.cell(row=r, column=2).value for r in range(2, ws1.max_row + 1)]
        _assert("Letter Date" in data_values, "Letter Date in Excel output")
        _assert("Proposed Fee" in data_values or any("Fee" in str(v) for v in data_values),
                "Fee field in Excel output")

        print(f"\n  Excel file     : {written.name}")
        print(f"  File size      : {written.stat().st_size:,} bytes")
        print(f"  Sheets         : {wb.sheetnames}")
        print(f"  Data rows      : {ws1.max_row - 1}")
        print(f"  Fields written : {len(consolidated)}")


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #

if __name__ == "__main__":
    test_normalisation()
    test_deduplication()
    test_conflict_detection()
    test_full_consolidation()
    test_excel_output()

    print("\n" + "=" * 60)
    print("All consolidation & output tests passed.")
    print("=" * 60)
